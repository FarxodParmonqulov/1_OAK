import logging
import re
import asyncio
import os
import sqlite3
import hashlib
import json
from docx import Document
from datetime import datetime, timedelta
from pathlib import Path
from dotenv import load_dotenv

# .env faylni yuklash (local development uchun)
load_dotenv()

# ================= CONFIG (ENV dan o'qiladi) =================
BOT_TOKEN = os.getenv("BOT_TOKEN", "")
GROUP_ID = int(os.getenv("GROUP_ID", "-1003600714782"))
GROUP_INVITE = os.getenv("GROUP_INVITE", "https://t.me/vacancy_argos_bugun")

# Railway da fayllar /app/ papkasida bo'ladi
OAK_FILE = os.getenv("OAK_FILE", "/app/oak_jurnallar.docx")
DOC_ROOT = os.getenv("DOC_ROOT", "/app/oak_docs")

# Railway Volume yoki /app/data papkasi
DB_FILE = os.getenv("DB_FILE", "/app/data/oak.db")
FILE_HASH_DB = os.getenv("FILE_HASH_DB", "/app/data/file_hashes.json")

CHECK_INTERVAL = int(os.getenv("CHECK_INTERVAL", "14400"))
ADMIN_IDS = [int(x) for x in os.getenv("ADMIN_IDS", "146270730").split(",")]

# Data papkasini yaratish
os.makedirs(os.path.dirname(DB_FILE), exist_ok=True)
os.makedirs(DOC_ROOT, exist_ok=True)

if not BOT_TOKEN:
    raise ValueError("❌ BOT_TOKEN env variable o'rnatilmagan!")

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(levelname)s - %(message)s"
)

from aiogram import Bot, Dispatcher, F
from aiogram.types import (
    Message, ReplyKeyboardMarkup, KeyboardButton,
    InlineKeyboardMarkup, InlineKeyboardButton, CallbackQuery
)
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import StatesGroup, State
from aiogram.fsm.storage.memory import MemoryStorage
from aiogram.filters import Command

bot = Bot(BOT_TOKEN)
dp = Dispatcher(storage=MemoryStorage())

# ================= FILE HASH MANAGEMENT =================
class FileHashManager:
    def __init__(self, hash_db_file=FILE_HASH_DB):
        self.hash_db_file = hash_db_file
        self.hashes = self._load_hashes()

    def _load_hashes(self):
        try:
            if os.path.exists(self.hash_db_file):
                with open(self.hash_db_file, 'r', encoding='utf-8') as f:
                    return json.load(f)
        except Exception as e:
            logging.error(f"Hash yuklashda xato: {e}")
        return {}

    def _save_hashes(self):
        try:
            os.makedirs(os.path.dirname(self.hash_db_file), exist_ok=True)
            with open(self.hash_db_file, 'w', encoding='utf-8') as f:
                json.dump(self.hashes, f, indent=2, ensure_ascii=False)
        except Exception as e:
            logging.error(f"Hash saqlashda xato: {e}")

    def get_file_hash(self, file_path):
        try:
            stat = os.stat(file_path)
            hash_data = f"{file_path}_{stat.st_mtime}_{stat.st_size}"
            return hashlib.md5(hash_data.encode()).hexdigest()
        except Exception as e:
            logging.error(f"Hash hisoblashda xato {file_path}: {e}")
            return None

    def is_file_changed(self, file_path):
        current_hash = self.get_file_hash(file_path)
        if not current_hash:
            return False
        try:
            file_key = str(Path(file_path).relative_to(DOC_ROOT))
        except ValueError:
            file_key = file_path
        if file_key not in self.hashes:
            self.hashes[file_key] = current_hash
            self._save_hashes()
            return True
        if self.hashes[file_key] != current_hash:
            self.hashes[file_key] = current_hash
            self._save_hashes()
            return True
        return False

    def remove_deleted_files(self, existing_files):
        current_files = set(self.hashes.keys())
        existing_set = set()
        for f in existing_files:
            try:
                existing_set.add(str(Path(f).relative_to(DOC_ROOT)))
            except ValueError:
                existing_set.add(f)
        deleted_files = current_files - existing_set
        for file_key in deleted_files:
            del self.hashes[file_key]
            logging.info(f"O'chirilgan fayl hash'dan olib tashlandi: {file_key}")
        if deleted_files:
            self._save_hashes()

file_hash_manager = FileHashManager()

# ================= DATABASE SETUP =================
def init_db():
    conn = sqlite3.connect(DB_FILE)
    cur = conn.cursor()
    cur.execute("""
    CREATE TABLE IF NOT EXISTS oak_docs (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        year TEXT,
        source TEXT,
        surname TEXT,
        fullname TEXT,
        text TEXT,
        additional_info TEXT,
        file_hash TEXT,
        notified INTEGER DEFAULT 0,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    )""")
    cur.execute("""
    CREATE TABLE IF NOT EXISTS user_watchlist (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        user_id INTEGER,
        fullname TEXT,
        surname TEXT,
        search_text TEXT,
        normalized_fio TEXT,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        UNIQUE(user_id, normalized_fio)
    )""")
    cur.execute("""
    CREATE TABLE IF NOT EXISTS sent_notifications (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        user_id INTEGER,
        doc_id INTEGER,
        source TEXT,
        fullname TEXT,
        sent_time DATETIME DEFAULT CURRENT_TIMESTAMP,
        UNIQUE(user_id, doc_id)
    )""")
    cur.execute("""
    CREATE TABLE IF NOT EXISTS oak_notify (
        user_id INTEGER,
        surname TEXT,
        fullname TEXT,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        UNIQUE(user_id, fullname)
    )""")
    cur.execute("""
    CREATE TABLE IF NOT EXISTS oak_notifications_sent (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        user_id INTEGER,
        source TEXT,
        fullname TEXT,
        sent_time DATETIME DEFAULT CURRENT_TIMESTAMP,
        UNIQUE(user_id, source, fullname)
    )""")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_oak_docs_year ON oak_docs(year)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_oak_docs_source ON oak_docs(source)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_oak_docs_fullname ON oak_docs(fullname)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_oak_docs_notified ON oak_docs(notified)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_oak_docs_surname ON oak_docs(surname)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_user_watchlist_user ON user_watchlist(user_id)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_user_watchlist_fio ON user_watchlist(normalized_fio)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_sent_notifications_user_doc ON sent_notifications(user_id, doc_id)")
    conn.commit()
    conn.close()
    logging.info("✅ Database tayyor")

init_db()

# ================= TEXT NORMALIZER =================
def normalize(text):
    if not text:
        return ""
    t = text.lower()
    t = t.replace("'","'").replace("'","'").replace("ʼ","'")
    t = t.replace("o'","o").replace("g'","g")
    t = t.replace("ў","o").replace("қ","q").replace("ғ","g").replace("ҳ","h")
    t = t.replace("o'g'li", "ўғли").replace("qizi", "қизи")
    t = re.sub(r"\s+"," ",t)
    return t.strip()

# ================= LATIN → CYRILLIC =================
LAT2CYR = {
    "o'":"ў","g'":"ғ","sh":"ш","ch":"ч","ya":"я","yo":"ё","yu":"ю",
    "a":"а","b":"б","d":"д","e":"е","f":"ф","g":"г","h":"х","i":"и",
    "j":"ж","k":"к","l":"л","m":"м","n":"н","o":"о","p":"п","q":"қ",
    "r":"р","s":"с","t":"т","u":"у","v":"в","x":"х","y":"й","z":"з",
    "o`":"ў","g`":"ғ","o'":"ў","g'":"ғ"
}

def latin_to_cyr(text):
    t = text.lower()
    for k,v in LAT2CYR.items():
        t = t.replace(k,v)
    return t

# ================= USER WATCHLIST =================
def add_to_watchlist(user_id, search_text, return_cyrillic=False):
    cyrillic_text = latin_to_cyr(search_text)
    normalized_fio = normalize(cyrillic_text)
    parts = normalized_fio.split()
    surname = parts[0] if parts else ""
    conn = sqlite3.connect(DB_FILE)
    cur = conn.cursor()
    try:
        cur.execute("""
            INSERT OR REPLACE INTO user_watchlist
            (user_id, fullname, surname, search_text, normalized_fio)
            VALUES (?, ?, ?, ?, ?)
        """, (user_id, cyrillic_text, surname, search_text, normalized_fio))
        conn.commit()
        logging.info(f"✅ Watchlistga qo'shildi: user_id={user_id}, fio={cyrillic_text}")
        if return_cyrillic:
            return True, cyrillic_text
        return True
    except Exception as e:
        logging.error(f"❌ Watchlistga qo'shishda xato: {e}")
        if return_cyrillic:
            return False, None
        return False
    finally:
        conn.close()

def get_user_watchlist(user_id):
    conn = sqlite3.connect(DB_FILE)
    cur = conn.cursor()
    cur.execute("""
        SELECT id, fullname, search_text, created_at
        FROM user_watchlist
        WHERE user_id = ?
        ORDER BY created_at DESC
    """, (user_id,))
    results = cur.fetchall()
    conn.close()
    return results

def remove_from_watchlist(user_id, watchlist_id):
    conn = sqlite3.connect(DB_FILE)
    cur = conn.cursor()
    cur.execute("DELETE FROM user_watchlist WHERE id = ? AND user_id = ?", (watchlist_id, user_id))
    deleted = cur.rowcount > 0
    conn.commit()
    conn.close()
    return deleted

# ================= EXACT MATCH =================
def find_exact_match_for_user(user_id):
    conn = sqlite3.connect(DB_FILE)
    cur = conn.cursor()
    cur.execute("SELECT normalized_fio, fullname, search_text FROM user_watchlist WHERE user_id = ?", (user_id,))
    user_watchlist = cur.fetchall()
    if not user_watchlist:
        conn.close()
        return []
    cur.execute("""
        SELECT id, year, source, fullname, text, additional_info
        FROM oak_docs WHERE notified = 0
        ORDER BY year DESC, source
    """)
    new_docs = cur.fetchall()
    matches = []
    for normalized_fio, cyrillic_fio, original_text in user_watchlist:
        search_variants = [
            normalized_fio,
            normalize(original_text),
            latin_to_cyr(original_text),
        ]
        for doc_id, year, source, doc_fullname, doc_text, doc_additional in new_docs:
            doc_normalized = normalize(doc_fullname)
            match_found = False
            if doc_normalized in search_variants:
                match_found = True
            elif normalize(latin_to_cyr(doc_fullname)) in search_variants:
                match_found = True
            elif (doc_normalized.replace("ўғли", "o'g'li") in search_variants or
                  doc_normalized.replace("o'g'li", "ўғли") in search_variants):
                match_found = True
            if match_found:
                cur.execute("""
                    SELECT COUNT(*) FROM sent_notifications
                    WHERE user_id = ? AND doc_id = ?
                """, (user_id, doc_id))
                already_sent = cur.fetchone()[0] > 0
                if not already_sent:
                    matches.append({
                        'doc_id': doc_id, 'year': year, 'source': source,
                        'doc_fullname': doc_fullname, 'doc_text': doc_text,
                        'doc_additional': doc_additional,
                        'user_search': original_text, 'user_cyrillic': cyrillic_fio
                    })
    conn.close()
    return matches

# ================= NOTIFICATION =================
async def send_single_notification(user_id, match_data):
    try:
        source = match_data['source']
        year = match_data['year']
        fullname = match_data['doc_fullname']
        text = match_data['doc_text']
        additional = match_data['doc_additional']
        match = re.match(r"(\d{4})-(\d)", source)
        if match:
            doc_year = match.group(1)
            doc_quarter = match.group(2)
            period = f"{doc_year}-yil, {doc_quarter}-chorak"
        else:
            period = source
        message = f"🎉 **OAK BYULLETEN - TOPILDI!**\n\n"
        message += f"👤 **{fullname}**\n"
        message += f"📅 {period}\n"
        if text and text != fullname:
            if len(text) > 200:
                text = text[:200] + "..."
            message += f"📝 {text}\n"
        if additional:
            if len(additional) > 150:
                additional = additional[:150] + "..."
            message += f"ℹ️ {additional}\n"
        await bot.send_message(user_id, message)
        conn = sqlite3.connect(DB_FILE)
        cur = conn.cursor()
        cur.execute("""
            INSERT OR IGNORE INTO sent_notifications (user_id, doc_id, source, fullname)
            VALUES (?, ?, ?, ?)
        """, (user_id, match_data['doc_id'], source, fullname))
        cur.execute("UPDATE oak_docs SET notified = 1 WHERE id = ?", (match_data['doc_id'],))
        conn.commit()
        conn.close()
        logging.info(f"📨 Xabar yuborildi: user_id={user_id}, doc_id={match_data['doc_id']}")
        return True
    except Exception as e:
        logging.error(f"❌ Xabar yuborishda xato user_id={user_id}: {e}")
        return False

async def check_all_users_after_scan():
    conn = sqlite3.connect(DB_FILE)
    cur = conn.cursor()
    cur.execute("SELECT DISTINCT user_id FROM user_watchlist")
    users = cur.fetchall()
    conn.close()
    if not users:
        return 0
    notified_count = 0
    for (user_id,) in users:
        try:
            matches = find_exact_match_for_user(user_id)
            if not matches:
                continue
            latest_match = max(matches, key=lambda x: f"{x['year']}{x['source']}")
            conn = sqlite3.connect(DB_FILE)
            cur = conn.cursor()
            cur.execute("""
                SELECT COUNT(*) FROM sent_notifications WHERE user_id = ? AND doc_id = ?
            """, (user_id, latest_match['doc_id']))
            already_sent = cur.fetchone()[0] > 0
            conn.close()
            if not already_sent:
                await send_single_notification(user_id, latest_match)
                notified_count += 1
                await asyncio.sleep(0.3)
        except Exception as e:
            logging.error(f"❌ Check xatosi user_id={user_id}: {e}")
    return notified_count

async def check_and_notify_user(user_id):
    matches = find_exact_match_for_user(user_id)
    if not matches:
        return False
    latest_match = max(matches, key=lambda x: f"{x['year']}{x['source']}")
    await send_single_notification(user_id, latest_match)
    return True

async def periodic_user_check():
    conn = sqlite3.connect(DB_FILE)
    cur = conn.cursor()
    cur.execute("SELECT DISTINCT user_id FROM user_watchlist")
    users = cur.fetchall()
    conn.close()
    notified_count = 0
    for (user_id,) in users:
        try:
            notified = await check_and_notify_user(user_id)
            if notified:
                notified_count += 1
            await asyncio.sleep(0.5)
        except Exception as e:
            logging.error(f"❌ User check xatosi user_id={user_id}: {e}")
    return notified_count

# ================= DOCX PARSER =================
def extract_people_from_doc(doc):
    lines = [p.text.replace("\u00a0"," ").strip() for p in doc.paragraphs if p.text.strip()]
    people = []
    organization_keywords = [
        "институти","университети","академияси","маркази","фонди",
        "қўмитаси","идораси","агентлиги","маҳкамаси","бўлими",
        "кафедраси","факультети","филиали","лабораторияси","тадқиқотчиси",
        "ўқитувчиси","донишгоҳи","шуъбаси","бошқармаси","назарияси",
        "муассасаси","илмий","тадқиқот"
    ]
    degree_codes = [
        "DSc.","PhD.","BSc.","MSc.","Dr.","Prof.",
        "01.01.01","02.00.09","02.00.13","02.00.14","03.00.06","05.02.03",
        "06.01.05","08.00.01","08.00.03","08.00.06","08.00.07","08.00.10",
        "13.00.02","14.00.13","14.00.40","19.00.04","19.00.05",
        "В2019","В2020","В2021","В2022","B2019","B2020","B2021","B2022",
    ]
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line or line.isspace():
            i += 1
            continue
        person_data = {"surname":"","fullname":"","text":"","additional_info":""}

        # FORMAT 1: "1512. Хожаметов Гулмурат Бекмуратович ..."
        if re.match(r'^\d+\.\s+[А-ЯЁЎҚҒҲ]', line) and ' ' in line and len(line) > 20:
            match = re.match(r'^\d+\.\s+(.+)$', line)
            if match:
                text_after_number = match.group(1)
                fio_match = re.match(
                    r'^([А-ЯЁЎҚҒҲ][а-яёўқғҳ]+\s+[А-ЯЁЎҚҒҲ][а-яёўқғҳ]+\s+[А-ЯЁЎҚҒҲ][а-яёўқғҳ]+(?:вич|вна|ўғли|қизи)?)\s+(.+)$',
                    text_after_number)
                if fio_match:
                    fullname = fio_match.group(1)
                    rest_of_line = fio_match.group(2)
                    name_parts = fullname.split()
                    if len(name_parts) >= 2:
                        person_data["surname"] = normalize(name_parts[0])
                        person_data["fullname"] = normalize(fullname)
                        person_data["text"] = fullname
                        if rest_of_line:
                            person_data["additional_info"] = rest_of_line[:300]
                        people.append(person_data)

        # FORMAT 2: Alohida qatorlar - Familiya / Ism / Otasining ismi
        elif re.match(r'^[А-ЯЁЎҚҒҲ][а-яёўқғҳ]+$', line) and i + 2 < len(lines):
            surname = line
            next_line1 = lines[i + 1]
            next_line2 = lines[i + 2]
            is_name = re.match(r'^[А-ЯЁЎҚҒҲ][а-яёўқғҳ]+$', next_line1)
            is_patronymic = re.match(r'^[А-ЯЁЎҚҒҲ][а-яёўқғҳ]+\s+(?:ўғли|қизи)$', next_line2) or \
                            re.match(r'^[А-ЯЁЎҚҒҲ][а-яёўқғҳ]+(?:вич|вна)$', next_line2)
            if is_name and is_patronymic:
                fullname = f"{surname} {next_line1} {next_line2}"
                person_data["surname"] = normalize(surname)
                person_data["fullname"] = normalize(fullname)
                person_data["text"] = fullname
                additional_info_parts = []
                for j in range(i + 3, min(i + 15, len(lines))):
                    nxt = lines[j]
                    if not nxt:
                        continue
                    if re.match(r'^[А-ЯЁЎҚҒҲ][а-яёўқғҳ]+$', nxt) and j < len(lines)-2:
                        if re.match(r'^[А-ЯЁЎҚҒҲ][а-яёўқғҳ]+$', lines[j+1]):
                            break
                    if any(code in nxt for code in degree_codes) or \
                       any(kw in nxt.lower() for kw in organization_keywords):
                        additional_info_parts.append(nxt)
                    elif len(nxt) < 150 and not re.match(r'^\d+\.\s*$', nxt):
                        additional_info_parts.append(nxt)
                if additional_info_parts:
                    person_data["additional_info"] = " | ".join(additional_info_parts[:5])
                people.append(person_data)
                i += 2

        # FORMAT 3: "АТАБЕКОВА КАМОЛА ГАЙРАТОВНА. «..."
        elif re.match(r'^[А-ЯЁЎҚҒҲ][А-ЯЁЎҚҒҲ\s]+\.\s+«', line) or \
             re.match(r'^[А-ЯЁЎҚҒҲ][А-ЯЁЎҚҒҲ\s]+\.\s*$', line):
            if '.' in line:
                parts = line.split('.', 1)
                name_part = parts[0].strip()
                name_parts = name_part.split()
                if len(name_parts) >= 2:
                    surname = name_parts[0].capitalize()
                    name = name_parts[1].capitalize()
                    patronymic = name_parts[2].capitalize() if len(name_parts) > 2 else ""
                    fullname = f"{surname} {name} {patronymic}".strip()
                    person_data["surname"] = normalize(surname)
                    person_data["fullname"] = normalize(fullname)
                    person_data["text"] = fullname
                    if len(parts) > 1 and parts[1].strip():
                        remaining = parts[1].strip()
                        person_data["additional_info"] = remaining[:250]
                    people.append(person_data)

        # FORMAT 4: FIO + raqamli kod
        elif re.match(r'^[А-ЯЁЎҚҒҲ][а-яёўқғҳ]+\s+[А-ЯЁЎҚҒҲ][а-яёўқғҳ]+\s+[А-ЯЁЎҚҒҲ][а-яёўқғҳ]+(?:вич|вна|ўғли|қизи)?\s+[\d\.]', line):
            match = re.match(
                r'^([А-ЯЁЎҚҒҲ][а-яёўқғҳ]+\s+[А-ЯЁЎҚҒҲ][а-яёўқғҳ]+\s+[А-ЯЁЎҚҒҲ][а-яёўқғҳ]+(?:вич|вна|ўғли|қизи)?)\s+(.+)$',
                line)
            if match:
                fullname = match.group(1)
                rest = match.group(2)
                name_parts = fullname.split()
                if len(name_parts) >= 2:
                    person_data["surname"] = normalize(name_parts[0])
                    person_data["fullname"] = normalize(fullname)
                    person_data["text"] = fullname
                    person_data["additional_info"] = rest[:300]
                    people.append(person_data)

        # FORMAT 7: KATTA HARFLAR "ТОРЕШОВА АМИНА УББИНИЯЗОВНА"
        elif re.match(r'^[А-ЯЁЎҚҒҲ]{2,}\s+[А-ЯЁЎҚҒҲ]{2,}\s+[А-ЯЁЎҚҒҲ]{2,}$', line):
            name_parts = line.split()
            if len(name_parts) >= 2:
                surname = name_parts[0].capitalize()
                name = name_parts[1].capitalize()
                patronymic = name_parts[2].capitalize() if len(name_parts) > 2 else ""
                fullname = f"{surname} {name} {patronymic}".strip()
                person_data["surname"] = normalize(surname)
                person_data["fullname"] = normalize(fullname)
                person_data["text"] = fullname
                additional_info_parts = []
                for j in range(i + 1, min(i + 5, len(lines))):
                    nxt = lines[j]
                    if not nxt:
                        continue
                    if re.match(r'^[А-ЯЁЎҚҒҲ]{2,}\s+[А-ЯЁЎҚҒҲ]', nxt):
                        break
                    if len(nxt) < 200:
                        additional_info_parts.append(nxt)
                if additional_info_parts:
                    person_data["additional_info"] = " | ".join(additional_info_parts[:3])
                people.append(person_data)

        # FORMAT 9: ўғли/қизи bilan
        elif re.match(r'^[А-ЯЁЎҚҒҲ][а-яёўқғҳ]+\s+[А-ЯЁЎҚҒҲ][а-яёўқғҳ]+\s+[А-ЯЁЎҚҒҲ][а-яёўқғҳ]+\s+(?:ўғли|қизи)\s+[\d\.]', line):
            match = re.match(
                r'^([А-ЯЁЎҚҒҲ][а-яёўқғҳ]+\s+[А-ЯЁЎҚҒҲ][а-яёўқғҳ]+\s+[А-ЯЁЎҚҒҲ][а-яёўқғҳ]+\s+(?:ўғли|қизи))\s+(.+)$',
                line)
            if match:
                fullname = match.group(1)
                rest = match.group(2)
                name_parts = fullname.split()
                if len(name_parts) >= 2:
                    person_data["surname"] = normalize(name_parts[0])
                    person_data["fullname"] = normalize(fullname)
                    person_data["text"] = fullname
                    person_data["additional_info"] = rest[:250]
                    people.append(person_data)

        # FORMAT 10: Raqam alohida "1512." keyin FIO
        elif re.match(r'^\d+\.\s*$', line) and i + 1 < len(lines):
            next_line = lines[i + 1]
            if re.match(r'^[А-ЯЁЎҚҒҲ][а-яёўқғҳ]+\s+[А-ЯЁЎҚҒҲ][а-яёўқғҳ]+\s+[А-ЯЁЎҚҒҲ][а-яёўқғҳ]+(?:вич|вна|ўғли|қизи)?', next_line):
                fullname = next_line
                name_parts = fullname.split()
                if len(name_parts) >= 2:
                    person_data["surname"] = normalize(name_parts[0])
                    person_data["fullname"] = normalize(fullname)
                    person_data["text"] = fullname
                    person_data["additional_info"] = f"№ {line.strip()}"
                    people.append(person_data)
                    i += 1

        i += 1

    # Takroriylarni olib tashlash
    unique_people = []
    seen_names = set()
    for person in people:
        if person['fullname'] and person['fullname'] not in seen_names:
            unique_people.append(person)
            seen_names.add(person['fullname'])
    logging.info(f"📊 Fayldan {len(unique_people)} ta odam topildi")
    return unique_people

# ================= FILE PROCESSING =================
async def process_new_or_changed_files(initial_scan=False):
    conn = sqlite3.connect(DB_FILE)
    cur = conn.cursor()
    all_files = []
    changed_files = []
    if not os.path.exists(DOC_ROOT):
        logging.warning(f"⚠️ DOC_ROOT papkasi topilmadi: {DOC_ROOT}")
        conn.close()
        return False
    for year_folder in os.listdir(DOC_ROOT):
        year_path = os.path.join(DOC_ROOT, year_folder)
        if not os.path.isdir(year_path):
            continue
        for filename in os.listdir(year_path):
            if filename.lower().endswith(".docx"):
                file_path = os.path.join(year_path, filename)
                all_files.append(file_path)
                if initial_scan:
                    changed_files.append((year_folder, filename, file_path))
                else:
                    if file_hash_manager.is_file_changed(file_path):
                        changed_files.append((year_folder, filename, file_path))
    file_hash_manager.remove_deleted_files(all_files)
    for year_folder, filename, file_path in changed_files:
        try:
            cur.execute("DELETE FROM oak_docs WHERE year=? AND source=?", (year_folder, filename))
            doc = Document(file_path)
            people = extract_people_from_doc(doc)
            file_hash = file_hash_manager.get_file_hash(file_path)
            if people:
                for p in people:
                    cur.execute("""
                        INSERT INTO oak_docs(year,source,surname,fullname,text,additional_info,file_hash,notified)
                        VALUES(?,?,?,?,?,?,?,0)
                    """, (year_folder, filename, p["surname"], p["fullname"], p["text"],
                          p.get("additional_info",""), file_hash))
                conn.commit()
                logging.info(f"✅ {year_folder}/{filename}: {len(people)} ta odam")
            else:
                logging.warning(f"⚠️ {year_folder}/{filename}: Hech qanday odam topilmadi")
        except Exception as e:
            logging.error(f"❌ Faylni qayta ishlashda xato {filename}: {e}")
            conn.rollback()
    conn.close()
    if changed_files:
        logging.info(f"🔍 {len(changed_files)} ta fayl qayta ishlandi")
        return True
    return False

# ================= SEARCH =================
def smart_name_match_all(year, user_text):
    user_normalized = normalize(user_text)
    user_cyrillic = latin_to_cyr(user_text)
    user_cyrillic_normalized = normalize(user_cyrillic)
    search_terms = [user_normalized]
    if user_cyrillic_normalized != user_normalized:
        search_terms.append(user_cyrillic_normalized)
    search_parts_set = set()
    for term in search_terms:
        parts = term.split()
        if len(parts) >= 2:
            search_parts_set.add((parts[0], parts[1]))
        elif len(parts) == 1:
            search_parts_set.add((parts[0], ""))
    conn = sqlite3.connect(DB_FILE)
    cur = conn.cursor()
    cur.execute("""
        SELECT source, fullname, text, additional_info
        FROM oak_docs WHERE year=? ORDER BY source
    """, (year,))
    rows = cur.fetchall()
    conn.close()
    exact_matches = []
    for src, fullname, text, additional_info in rows:
        if not fullname:
            continue
        fn_normalized = normalize(fullname)
        fn_parts = fn_normalized.split()
        if len(fn_parts) >= 2:
            db_surname = fn_parts[0]
            db_name = fn_parts[1]
            for search_surname, search_name in search_parts_set:
                if search_surname and db_surname != search_surname:
                    continue
                if search_name and db_name != search_name:
                    continue
                full_text = text
                if additional_info:
                    full_text += f"\n📝 {additional_info}"
                exact_matches.append({
                    "source": src, "fullname": fullname,
                    "text": full_text,
                    "score": 10 if search_name else 5
                })
                break
    exact_matches.sort(key=lambda x: x["score"], reverse=True)
    return exact_matches

def format_oak_source(filename):
    name = os.path.splitext(filename)[0]
    m = re.match(r"(\d{4})-(\d)", name)
    if not m:
        return filename
    return f"📅 {m.group(1)}-yil, {m.group(2)}-chorak"

# ================= OAK JOURNALS =================
def load_oak():
    if not os.path.exists(OAK_FILE):
        logging.warning(f"⚠️ OAK fayli topilmadi: {OAK_FILE}")
        return []
    try:
        doc = Document(OAK_FILE)
        lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        blocks, buf = [], []
        for line in lines:
            if "OAK Rayosatining" in line and buf:
                blocks.append(buf)
                buf = []
            buf.append(line)
        if buf:
            blocks.append(buf)
        return blocks
    except Exception as e:
        logging.error(f"❌ OAK fayl yuklashda xato: {e}")
        return []

OAK_BLOCKS = load_oak()

def find_journal(name):
    name = normalize(name)
    results = []
    for block in OAK_BLOCKS:
        qaror = block[0]
        i = 1
        while i < len(block):
            line = block[i]
            if re.match(r"\d{2}\.\d{2}\.\d{2}\s*–", line):
                current_ixt = line
                i += 1
                buf = []
                while i < len(block) and not re.match(r"\d{2}\.\d{2}\.\d{2}\s*–", block[i]):
                    buf.append(block[i])
                    i += 1
                text = normalize(" ".join(buf))
                if name in text:
                    results.append(
                        "📜 OAK qarori:\n" + qaror + "\n\n"
                        "🎓 Ixtisoslik:\n" + current_ixt + "\n\n"
                        "📘 Jurnal ma'lumotlari:\n" + "\n".join(buf)
                    )
            else:
                i += 1
    if results:
        return "\n\n━━━━━━━━━━━━━━━━━━━━\n\n".join(results)
    return None

# ================= KEYBOARDS =================
main_menu = ReplyKeyboardMarkup(
    keyboard=[
        [KeyboardButton(text="📘 OAK jurnallarini tekshirish")],
        [KeyboardButton(text="🔍 Byulletendan mavzuni qidirish")],
        [KeyboardButton(text="📢 OAK bildirishnoma")],
        [KeyboardButton(text="📋 Mening kuzatish ro'yxatim")],
    ],
    resize_keyboard=True
)

def year_keyboard():
    rows, row = [], []
    for y in range(2016, 2027):
        row.append(KeyboardButton(text=str(y)))
        if len(row) == 4:
            rows.append(row)
            row = []
    if row:
        rows.append(row)
    rows.append([KeyboardButton(text="⬅️ Orqaga")])
    return ReplyKeyboardMarkup(keyboard=rows, resize_keyboard=True)

# ================= GROUP CHECK =================
async def is_member(user_id):
    try:
        m = await bot.get_chat_member(GROUP_ID, user_id)
        return m.status in ["member","administrator","creator"]
    except:
        return False

# ================= HANDLERS =================
@dp.message(Command("start"))
async def start(m: Message):
    if await is_member(m.from_user.id):
        await m.answer("✅ Bot tayyor", reply_markup=main_menu)
    else:
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="➕ Guruhga kirish", url=GROUP_INVITE)],
            [InlineKeyboardButton(text="✅ Tekshirish", callback_data="check")]
        ])
        await m.answer("🔐 Botdan foydalanish uchun guruhga a'zo bo'ling:", reply_markup=kb)

@dp.callback_query(F.data == "check")
async def check(call: CallbackQuery):
    if await is_member(call.from_user.id):
        await call.message.delete()
        await call.message.answer("🎉 Tasdiqlandi!\n\n📋 Asosiy menyu:", reply_markup=main_menu)
    else:
        await call.answer("❌ Hali guruhga a'zo emassiz", show_alert=True)

# OAK Jurnallar
class JournalSearch(StatesGroup):
    waiting = State()

@dp.message(F.text == "📘 OAK jurnallarini tekshirish")
async def ask_journal(m: Message, state: FSMContext):
    await m.answer("📘 Jurnal nomini kiriting:",
                   reply_markup=ReplyKeyboardMarkup(
                       keyboard=[[KeyboardButton(text="⬅️ Orqaga")]],
                       resize_keyboard=True))
    await state.set_state(JournalSearch.waiting)

@dp.message(JournalSearch.waiting)
async def search_journal(m: Message, state: FSMContext):
    if m.text == "⬅️ Orqaga":
        await m.answer("📋 Asosiy menyu:", reply_markup=main_menu)
        await state.clear()
        return
    res = find_journal(m.text)
    if res:
        await m.answer("✅ Jurnal OAK ro'yxatida topildi:\n\n" + res)
    else:
        await m.answer("❌ Ushbu jurnal OAK ro'yxatida topilmadi.")
    await state.clear()
    await m.answer("📋 Menyu:", reply_markup=main_menu)

# Byulleten qidirish
class TopicSearch(StatesGroup):
    year = State()
    query = State()

@dp.message(F.text == "🔍 Byulletendan mavzuni qidirish")
async def choose_year(m: Message, state: FSMContext):
    await m.answer("📅 Yilni tanlang:", reply_markup=year_keyboard())
    await state.set_state(TopicSearch.year)

@dp.message(TopicSearch.year)
async def set_year(m: Message, state: FSMContext):
    if m.text == "⬅️ Orqaga":
        await m.answer("📋 Asosiy menyu:", reply_markup=main_menu)
        await state.clear()
        return
    await state.update_data(year=m.text)
    await m.answer("🔎 To'liq FIO kiriting:\n\nMisol: Хожаметов Гулмурат Бекмуратович\nyoki: Xojametov Gulmurat Bekmuratovich",
                   reply_markup=ReplyKeyboardMarkup(
                       keyboard=[[KeyboardButton(text="⬅️ Orqaga")]],
                       resize_keyboard=True))
    await state.set_state(TopicSearch.query)

def index_year_if_needed(year):
    conn = sqlite3.connect(DB_FILE)
    cur = conn.cursor()
    cur.execute("SELECT COUNT(*) FROM oak_docs WHERE year=?", (year,))
    if cur.fetchone()[0] > 0:
        conn.close()
        return
    folder = os.path.join(DOC_ROOT, year)
    if not os.path.isdir(folder):
        conn.close()
        return
    for f in sorted(os.listdir(folder)):
        if f.lower().endswith(".docx"):
            try:
                doc = Document(os.path.join(folder, f))
                people = extract_people_from_doc(doc)
                for p in people:
                    cur.execute("""
                        INSERT INTO oak_docs(year,source,surname,fullname,text,additional_info)
                        VALUES(?,?,?,?,?,?)
                    """, (year, f, p["surname"], p["fullname"], p["text"], p.get("additional_info","")))
            except Exception as e:
                logging.error(f"DOCX xato: {f} {e}")
    conn.commit()
    conn.close()

@dp.message(TopicSearch.query)
async def do_search(m: Message, state: FSMContext):
    if m.text == "⬅️ Orqaga":
        await m.answer("📅 Yilni tanlang:", reply_markup=year_keyboard())
        await state.set_state(TopicSearch.year)
        return
    user_id = m.from_user.id
    search_text = m.text.strip()
    year = (await state.get_data()).get('year', '2024')
    index_year_if_needed(year)
    matches = smart_name_match_all(year, search_text)
    if not matches:
        await m.answer(f"❌ {year}-yilda '{search_text}' uchun natija topilmadi.",
                       reply_markup=main_menu)
        await state.clear()
        return
    response = f"✅ {year}-yilda {len(matches)} ta natija topildi:\n\n"
    for i, match in enumerate(matches[:10], 1):
        source = match["source"]
        fullname = match["fullname"]
        text = match["text"]
        response += f"{i}. 👤 **{fullname}**\n"
        response += f"   📂 {format_oak_source(source)}\n"
        if text and text != fullname:
            if len(text) > 100:
                text = text[:100] + "..."
            response += f"   📝 {text}\n"
        response += "\n"
    if len(matches) > 10:
        response += f"\n... va yana {len(matches)-10} ta natija.\n"
    await m.answer(response, parse_mode="Markdown")
    await m.answer("📢 **Kuzatish ro'yxatiga qo'shilsinmi?**",
                   reply_markup=InlineKeyboardMarkup(inline_keyboard=[[
                       InlineKeyboardButton(text="✅ Ha", callback_data=f"add_watch:{search_text}"),
                       InlineKeyboardButton(text="❌ Yo'q", callback_data="cancel_to_main")
                   ]]))
    await state.clear()

# Callbacks
@dp.callback_query(F.data.startswith("add_watch:"))
async def add_to_watch_callback(call: CallbackQuery):
    search_text = call.data.split(":", 1)[1]
    success, cyrillic_fio = add_to_watchlist(call.from_user.id, search_text, return_cyrillic=True)
    if success:
        await call.message.edit_text(
            f"✅ **Kuzatish ro'yxatiga qo'shildi!**\n\n"
            f"🔍 Qidirilgan: {search_text}\n"
            f"🔤 Krilcha: {cyrillic_fio}\n\n"
            f"Yangi byulletenlarda topilsa xabar beriladi (har 4 soatda).",
            reply_markup=None)
    else:
        await call.message.edit_text("❌ Xatolik yuz berdi.", reply_markup=None)
    await call.answer()
    await call.message.answer("📋 Asosiy menyu:", reply_markup=main_menu)

@dp.callback_query(F.data == "cancel_to_main")
async def cancel_to_main_callback(call: CallbackQuery):
    await call.message.edit_text("❌ Kuzatish ro'yxatiga qo'shilmadi.", reply_markup=None)
    await call.answer()
    await call.message.answer("📋 Asosiy menyu:", reply_markup=main_menu)

@dp.callback_query(F.data == "cancel_watch")
async def cancel_watch_callback(call: CallbackQuery):
    await call.message.edit_text("❌ Kuzatish ro'yxatiga qo'shilmadi.", reply_markup=None)
    await call.answer()
    await call.message.answer("📋 Asosiy menyu:", reply_markup=main_menu)

# Watchlist
@dp.message(F.text == "📋 Mening kuzatish ro'yxatim")
async def my_watchlist(m: Message):
    watchlist = get_user_watchlist(m.from_user.id)
    if not watchlist:
        await m.answer("📭 Sizning kuzatish ro'yxatingiz bo'sh.\n\n"
                       "🔍 Byulletendan mavzuni qidirish bo'limida qidiring.",
                       reply_markup=main_menu)
        return
    response = "📋 **Sizning kuzatish ro'yxatingiz:**\n\n"
    for i, (watch_id, cyrillic_fio, original_text, created_at) in enumerate(watchlist, 1):
        response += f"{i}. **{cyrillic_fio}**\n"
        response += f"   📝 {original_text}\n"
        response += f"   📅 {created_at[:10]}\n\n"
    keyboard = []
    for watch_id, cyrillic_fio, _, _ in watchlist[:5]:
        keyboard.append([InlineKeyboardButton(
            text=f"❌ {cyrillic_fio[:20]}",
            callback_data=f"remove_watch:{watch_id}"
        )])
    keyboard.append([InlineKeyboardButton(text="⬅️ Asosiy menyu", callback_data="back_to_main")])
    await m.answer(response, reply_markup=InlineKeyboardMarkup(inline_keyboard=keyboard), parse_mode="Markdown")

@dp.callback_query(F.data.startswith("remove_watch:"))
async def remove_watch_callback(call: CallbackQuery):
    watch_id = int(call.data.split(":", 1)[1])
    deleted = remove_from_watchlist(call.from_user.id, watch_id)
    await call.message.edit_text(
        "✅ O'chirildi." if deleted else "❌ O'chirishda xatolik.", reply_markup=None)
    await call.answer()
    await call.message.answer("📋 Asosiy menyu:", reply_markup=main_menu)

@dp.callback_query(F.data == "back_to_main")
async def back_to_main_callback(call: CallbackQuery):
    await call.message.edit_text("📋 Asosiy menyuga qaytildi.", reply_markup=None)
    await call.answer()
    await call.message.answer("📋 Asosiy menyu:", reply_markup=main_menu)

# OAK bildirishnoma
class NotifyState(StatesGroup):
    waiting_fio = State()

@dp.message(F.text == "📢 OAK bildirishnoma")
async def notify_start(m: Message, state: FSMContext):
    watchlist = get_user_watchlist(m.from_user.id)
    if watchlist:
        response = "📋 **Hozirgi kuzatish ro'yxatingiz:**\n\n"
        for i, (_, cyrillic_fio, original_text, _) in enumerate(watchlist, 1):
            response += f"{i}. **{cyrillic_fio}**\n   📝 {original_text}\n\n"
        await m.answer(response, parse_mode="Markdown")
    await m.answer(
        "📢 Yangi FIO ni kuzatish ro'yxatiga qo'shish:\n\nMisol: Хожаметов Гулмурат Бекмуратович",
        reply_markup=ReplyKeyboardMarkup(
            keyboard=[[KeyboardButton(text="⬅️ Orqaga")]],
            resize_keyboard=True))
    await state.set_state(NotifyState.waiting_fio)

@dp.message(NotifyState.waiting_fio)
async def save_notify(m: Message, state: FSMContext):
    if m.text == "⬅️ Orqaga":
        await m.answer("📋 Asosiy menyu:", reply_markup=main_menu)
        await state.clear()
        return
    success, cyrillic_fio = add_to_watchlist(m.from_user.id, m.text.strip(), return_cyrillic=True)
    if success:
        await m.answer(
            f"✅ **Kuzatish ro'yxatiga qo'shildi!**\n\n"
            f"🔍 Qidirilgan: {m.text.strip()}\n"
            f"🔤 Krilcha: {cyrillic_fio}\n\n"
            f"Yangi byulletenlarda topilsa xabar beriladi.",
            reply_markup=main_menu)
    else:
        await m.answer("❌ Xatolik yuz berdi.", reply_markup=main_menu)
    await state.clear()

# Admin
@dp.message(Command("admin"))
async def admin_panel(m: Message):
    if m.from_user.id not in ADMIN_IDS:
        await m.answer("❌ Siz admin emassiz")
        return
    admin_menu = ReplyKeyboardMarkup(
        keyboard=[
            [KeyboardButton(text="📊 Statistika")],
            [KeyboardButton(text="🔄 Fayllarni qayta skanerlash")],
            [KeyboardButton(text="🧹 Bazani tozalash")],
            [KeyboardButton(text="🔍 Bazani tekshirish")],
            [KeyboardButton(text="📨 Xabarlar tarixi")],
            [KeyboardButton(text="👥 Kuzatish ro'yxati statistika")],
            [KeyboardButton(text="⬅️ Asosiy menyu")]
        ], resize_keyboard=True)
    await m.answer("🔧 Admin panel:", reply_markup=admin_menu)

@dp.message(F.text == "📊 Statistika")
async def show_stats(m: Message):
    conn = sqlite3.connect(DB_FILE)
    cur = conn.cursor()
    cur.execute("SELECT COUNT(*) FROM oak_docs")
    total_docs = cur.fetchone()[0]
    cur.execute("SELECT COUNT(DISTINCT source) FROM oak_docs")
    total_files = cur.fetchone()[0]
    cur.execute("SELECT COUNT(*) FROM user_watchlist")
    total_watchlist = cur.fetchone()[0]
    cur.execute("SELECT COUNT(DISTINCT user_id) FROM user_watchlist")
    unique_users = cur.fetchone()[0]
    cur.execute("SELECT COUNT(*) FROM sent_notifications")
    total_notifications = cur.fetchone()[0]
    cur.execute("""
        SELECT year, COUNT(DISTINCT source), COUNT(*)
        FROM oak_docs GROUP BY year ORDER BY year
    """)
    year_stats = cur.fetchall()
    conn.close()
    stats_text = (
        f"📊 **Bot statistika:**\n\n"
        f"• Jami yozuvlar: {total_docs}\n"
        f"• Fayllar soni: {total_files}\n"
        f"• Kuzatish ro'yxati: {total_watchlist}\n"
        f"• Faol foydalanuvchilar: {unique_users}\n"
        f"• Yuborilgan xabarlar: {total_notifications}\n\n"
        f"**Yillar bo'yicha:**\n"
    )
    for year, file_count, entry_count in year_stats:
        stats_text += f"• {year}: {file_count} fayl, {entry_count} yozuv\n"
    await m.answer(stats_text)

@dp.message(F.text == "👥 Kuzatish ro'yxati statistika")
async def watchlist_stats(m: Message):
    conn = sqlite3.connect(DB_FILE)
    cur = conn.cursor()
    cur.execute("SELECT COUNT(*) FROM user_watchlist")
    total = cur.fetchone()[0]
    cur.execute("SELECT COUNT(DISTINCT user_id) FROM user_watchlist")
    users = cur.fetchone()[0]
    cur.execute("SELECT COUNT(*) FROM sent_notifications")
    sent = cur.fetchone()[0]
    cur.execute("""
        SELECT normalized_fio, COUNT(*) as count FROM user_watchlist
        GROUP BY normalized_fio ORDER BY count DESC LIMIT 5
    """)
    top_fios = cur.fetchall()
    conn.close()
    text = (f"📊 **Kuzatish ro'yxati:**\n\n"
            f"• Jami: {total}\n• Foydalanuvchilar: {users}\n• Yuborilgan: {sent}\n\n"
            f"**Top 5 FIO:**\n")
    for fio, count in top_fios:
        text += f"• {fio[:30]}: {count} ta\n"
    await m.answer(text)

@dp.message(F.text == "🔄 Fayllarni qayta skanerlash")
async def force_rescan(m: Message):
    await m.answer("🔄 Qayta skanerlash boshlandi...")
    global file_hash_manager
    file_hash_manager.hashes = {}
    file_hash_manager._save_hashes()
    has_changes = await process_new_or_changed_files(initial_scan=True)
    if has_changes:
        notified_count = await check_all_users_after_scan()
        await m.answer(f"✅ Muvaffaqiyatli. {notified_count} ta foydalanuvchiga xabar yuborildi.")
    else:
        await m.answer("ℹ️ O'zgarishlar topilmadi")

@dp.message(F.text == "🧹 Bazani tozalash")
async def cleanup_database(m: Message):
    conn = sqlite3.connect(DB_FILE)
    cur = conn.cursor()
    all_files_in_fs = []
    if os.path.exists(DOC_ROOT):
        for year_folder in os.listdir(DOC_ROOT):
            year_path = os.path.join(DOC_ROOT, year_folder)
            if os.path.isdir(year_path):
                for filename in os.listdir(year_path):
                    if filename.lower().endswith(".docx"):
                        all_files_in_fs.append((year_folder, filename))
    cur.execute("SELECT DISTINCT year, source FROM oak_docs")
    db_files = cur.fetchall()
    deleted_count = 0
    for year, source in db_files:
        if (year, source) not in all_files_in_fs:
            cur.execute("DELETE FROM oak_docs WHERE year=? AND source=?", (year, source))
            deleted_count += 1
    conn.commit()
    conn.close()
    await m.answer(f"🧹 {deleted_count} ta o'chirilgan fayl bazadan tozalandi")

@dp.message(F.text == "🔍 Bazani tekshirish")
async def check_database(m: Message):
    conn = sqlite3.connect(DB_FILE)
    cur = conn.cursor()
    cur.execute("""
        SELECT year, source, fullname, text, additional_info
        FROM oak_docs ORDER BY year DESC, source LIMIT 10
    """)
    entries = cur.fetchall()
    conn.close()
    if not entries:
        await m.answer("📭 Bazada hech qanday yozuv yo'q")
        return
    text = "🔍 **Bazadagi yozuvlar (oxirgi 10):**\n\n"
    for i, (year, source, fullname, t, add) in enumerate(entries, 1):
        text += f"{i}. {year}/{source}\n   👤 {fullname}\n"
        if t and t != fullname:
            text += f"   📝 {t[:80]}...\n" if len(t) > 80 else f"   📝 {t}\n"
        text += "\n"
    await m.answer(text)

@dp.message(F.text == "📨 Xabarlar tarixi")
async def notification_history(m: Message):
    conn = sqlite3.connect(DB_FILE)
    cur = conn.cursor()
    try:
        cur.execute("""
            SELECT sn.user_id, sn.source, sn.fullname, sn.sent_time
            FROM sent_notifications sn
            ORDER BY sn.sent_time DESC LIMIT 10
        """)
        notifications = cur.fetchall()
    except Exception as e:
        notifications = []
        logging.error(f"Notification history xato: {e}")
    conn.close()
    if not notifications:
        await m.answer("📭 Hozircha xabar yuborilmagan")
        return
    history_text = "📨 **Oxirgi xabarlar:**\n\n"
    for i, (user_id, source, fullname, sent_time) in enumerate(notifications, 1):
        history_text += f"{i}. 👤 {fullname}\n   📂 {source}\n   ⏰ {sent_time}\n\n"
    await m.answer(history_text)

@dp.message(F.text == "⬅️ Asosiy menyu")
async def back_to_main(m: Message):
    await m.answer("📋 Asosiy menyu:", reply_markup=main_menu)

# ================= PERIODIC CHECK =================
async def optimized_periodic_check():
    logging.info("🔄 Dastlabki skanerlash boshlandi...")
    has_changes = await process_new_or_changed_files(initial_scan=True)
    logging.info("✅ Dastlabki skanerlash yakunlandi")
    if has_changes:
        await check_all_users_after_scan()
    last_check_time = datetime.now()
    while True:
        try:
            current_time = datetime.now()
            elapsed = (current_time - last_check_time).total_seconds()
            if elapsed < CHECK_INTERVAL:
                await asyncio.sleep(600)
                continue
            logging.info(f"🔄 Tekshiruv boshlandi ({current_time})...")
            has_changes = await process_new_or_changed_files()
            if has_changes:
                notified_count = await periodic_user_check()
                logging.info(f"✅ {notified_count} ta foydalanuvchi xabardor qilindi")
            last_check_time = current_time
            next_check = last_check_time + timedelta(seconds=CHECK_INTERVAL)
            logging.info(f"✅ Keyingi tekshiruv: {next_check}")
        except Exception as e:
            logging.error(f"❌ Periodic check xatosi: {e}")
            await asyncio.sleep(60)

# ================= MAIN =================
async def main():
    logging.info("🚀 Bot ishga tushmoqda...")
    try:
        me = await bot.get_me()
        logging.info(f"✅ Bot ulandi: @{me.username}")
        await bot.delete_webhook(drop_pending_updates=True)
        asyncio.create_task(optimized_periodic_check())
        logging.info("🤖 Polling boshlandi...")
        await dp.start_polling(bot, skip_updates=True)
    except Exception as e:
        logging.error(f"❌ Xatolik: {e}")
        import traceback
        logging.error(traceback.format_exc())

if __name__ == "__main__":
    try:
        asyncio.run(main())
    except KeyboardInterrupt:
        logging.info("👋 Bot to'xtatildi")
    except Exception as e:
        logging.error(f"❌ Asosiy xatolik: {e}")
